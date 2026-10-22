#!/usr/bin/env python3
"""Proof Center — contre-calcul indépendant du CFO Cockpit.

Relit l'Excel cellule par cellule avec openpyxl (sans pandas, sans importer build_metrics.py),
recalcule les agrégats, les confronte à l'Excel (formules) et à metrics.json (chiffres affichés),
puis enquête sur chaque écart. Écrit proof.json et l'injecte dans cockpit.html.

Usage : python3 build_metrics.py && python3 verify.py
"""
import itertools
import json
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent
XLSX = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data" / "CFO_Cockpit_Delices_de_Bourgogne_2026.xlsx"
METRICS = ROOT / "metrics.json"
OUT = ROOT / "proof.json"
HTML = ROOT / "cockpit.html"
SEUIL = 1.0  # vert si |écart| < 1 €
NB = 9       # janv.–sept.
LIB = ["Janv", "Févr", "Mars", "Avr", "Mai", "Juin", "Juil", "Août", "Sept"]

wb = openpyxl.load_workbook(XLSX, data_only=False)   # saisies brutes (les formules restent des chaînes)
wbc = openpyxl.load_workbook(XLSX, data_only=True)   # valeurs calculées par Excel (cache)
M = json.loads(METRICS.read_text(encoding="utf-8"))


def e(x):
    return f"{x:,.2f}".replace(",", " ").replace(".", ",") + " €"


def lignes_par_libelle(ws):
    return {str(r[0].value).strip(): r[0].row for r in ws.iter_rows(min_row=2, max_col=1) if r[0].value}


def mois_saisis(ws, libelle, idx=None):
    """Valeurs saisies janv.–sept. d'une ligne ; refuse les formules (on ne veut que du brut)."""
    row = (idx or lignes_par_libelle(ws))[libelle]
    out = []
    for c in range(2, 2 + NB):
        v = ws.cell(row, c).value
        if isinstance(v, str) and v.startswith("="):
            raise ValueError(f"{ws.title}!{libelle} col {c} est une formule, pas une saisie")
        out.append(float(v or 0))
    return out


def mois_cache(nom, libelle):
    ws = wbc[nom]
    row = lignes_par_libelle(ws)[libelle]
    return [None if ws.cell(row, c).value is None else float(ws.cell(row, c).value) for c in range(2, 2 + NB)]


# ------------------------------------------------------------------ lecture brute
def pl_brut(nom):
    ws = wb[nom]
    idx = lignes_par_libelle(ws)
    gam = [k for k in idx if k.startswith("CA ")]
    ca_g = {k[3:]: mois_saisis(ws, k, idx) for k in gam}
    ca = [sum(ca_g[g][i] for g in ca_g) for i in range(NB)]
    cogs = mois_saisis(ws, "Coût des ventes", idx)
    postes = ["Frais de personnel", "Loyers & charges locatives", "Transport & logistique", "Marketing & commercial", "Autres charges externes"]
    ch = {p: mois_saisis(ws, p, idx) for p in postes}
    ebitda = [ca[i] - cogs[i] - sum(ch[p][i] for p in postes) for i in range(NB)]
    return {"ca_g": ca_g, "ca": ca, "cogs": cogs, "charges": ch, "ebitda": ebitda}


reel, bud = pl_brut("Reel_2026"), pl_brut("Budget_2026")
GAMMES = list(reel["ca_g"])

ws = wb["Ventes_Detail"]
hdr = [c.value for c in ws[1]]
col = {h: i for i, h in enumerate(hdr)}
ventes = []
for r in ws.iter_rows(min_row=2, values_only=True):
    if r[0] is None:
        continue
    ventes.append({"date": r[col["Date"]], "num": r[col["N° facture"]], "client": r[col["Client"]], "gamme": r[col["Gamme"]],
                   "produit": r[col["Produit"]], "qte": r[col["Quantité"]], "pu": r[col["Prix unitaire HT"]],
                   "ca": float(r[col["CA HT"]]), "cout": float(r[col["Coût de revient"]]), "ligne": None})
for i, v in enumerate(ventes):
    v["ligne"] = i + 2
    v["m"] = v["date"].month

hyp = {str(r[0]).strip(): r[1] for r in wb["Hypotheses"].iter_rows(min_row=2, values_only=True) if r[0]}
tva_alim, tva_alcool = float(hyp["Taux de TVA produits alimentaires"]), float(hyp["Taux de TVA alcools"])
taux_tva = lambda g: tva_alcool if "cassis" in g.lower() else tva_alim

cr = []
for r in wb["Creances_Clients"].iter_rows(min_row=2, values_only=True):
    if r[0] is None:
        continue
    cr.append({"num": r[0], "client": r[1], "date": r[2], "echeance": r[3], "ttc": float(r[4])})

tws = wb["Tresorerie"]
tidx = lignes_par_libelle(tws)
t_debut_janv = float(tws.cell(tidx["Trésorerie début de mois"], 2).value)
t_enc = mois_saisis(tws, "Encaissements clients", tidx)
SORTIES = ["Décaissements fournisseurs", "Salaires nets", "Charges sociales", "TVA décaissée", "Acomptes d'IS", "Loyers", "Investissements", "Remboursement d'emprunt"]
t_sort = {k: mois_saisis(tws, k, tidx) for k in SORTIES}

# ------------------------------------------------------------------ agrégats indépendants
det_ca = [0.0] * NB
det_cout = [0.0] * NB
det_g = defaultdict(lambda: [0.0] * NB)
det_ttc = [0.0] * NB
for v in ventes:
    i = v["m"] - 1
    det_ca[i] += v["ca"]
    det_cout[i] += v["cout"]
    det_g[v["gamme"]][i] += v["ca"]
    det_ttc[i] += v["ca"] * (1 + taux_tva(v["gamme"]))

treso_fin, t = [], t_debut_janv
for i in range(NB):
    t = t + t_enc[i] - sum(t_sort[k][i] for k in SORTIES)
    treso_fin.append(t)

cr_mois = [0.0] * NB
for c in cr:
    if c["date"].year == 2026 and 1 <= c["date"].month <= NB:
        cr_mois[c["date"].month - 1] += c["ttc"]

# ------------------------------------------------------------------ matrice de contrôle
lignes = []


def ligne(id_, controle, cat, la, lb, a, b, note=None, ecart_fn=None):
    cells = []
    for i in range(NB):
        if a[i] is None or b[i] is None:
            cells.append({"a": a[i], "b": b[i], "ecart": None, "ok": None})
            continue
        ec = ecart_fn(a[i], b[i]) if ecart_fn else a[i] - b[i]
        cells.append({"a": round(a[i], 2), "b": round(b[i], 2), "ecart": round(ec, 2), "ok": abs(ec) < SEUIL})
    ta = sum(x for x in a if x is not None) if all(x is not None for x in a) else None
    tb = sum(x for x in b if x is not None) if all(x is not None for x in b) else None
    tot = None if ta is None or tb is None else {"a": round(ta, 2), "b": round(tb, 2),
                                                  "ecart": round(ecart_fn(ta, tb) if ecart_fn else ta - tb, 2)}
    if tot:
        tot["ok"] = abs(tot["ecart"]) < SEUIL
    lignes.append({"id": id_, "controle": controle, "categorie": cat, "a_label": la, "b_label": lb,
                   "cellules": cells, "total": tot, "note": note})


ligne("ca", "CA détail des ventes vs CA compte de résultat", "Chiffre d'affaires", "Ventes_Detail (Σ CA HT)", "Reel_2026 (Σ CA gammes)", det_ca, reel["ca"])
for g in GAMMES:
    ligne("ca_" + g, f"↳ {g}", "Chiffre d'affaires", "Ventes_Detail", "Reel_2026", det_g[g], reel["ca_g"][g])
ligne("cogs", "Coût de revient détail vs coût des ventes", "Chiffre d'affaires", "Ventes_Detail (Σ coût)", "Reel_2026 (coût des ventes)", det_cout, reel["cogs"])
ligne("ca_aff", "CA recalculé vs CA affiché", "Cockpit", "Recalcul Excel brut", "metrics.json", reel["ca"], M["kpis"]["ca"]["mensuel"]["reel"])
ligne("ebitda_aff", "EBITDA recalculé vs EBITDA affiché", "Cockpit", "Recalcul Excel brut", "metrics.json", reel["ebitda"], M["kpis"]["ebitda"]["mensuel"]["reel"])
ligne("ebitda_bud_aff", "EBITDA budget recalculé vs affiché", "Cockpit", "Recalcul Excel brut", "metrics.json", bud["ebitda"], M["kpis"]["ebitda"]["mensuel"]["budget"])
ligne("ebitda_xl", "EBITDA recalculé vs formule Excel", "Excel", "Recalcul Excel brut", "Reel_2026 (cache formule)", reel["ebitda"], mois_cache("Reel_2026", "EBITDA"))
ligne("treso_xl", "Trésorerie fin de mois recalculée vs Excel", "Trésorerie", "Recalcul (début + flux)", "Tresorerie (cache formule)", treso_fin, mois_cache("Tresorerie", "Trésorerie fin de mois"))
ligne("treso_aff", "Trésorerie fin de mois recalculée vs affichée", "Trésorerie", "Recalcul (début + flux)", "metrics.json",
      treso_fin, [None] * (NB - 1) + [M["simulation"]["modele"]["treso"]["debut"]], note="Le cockpit n'affiche que la trésorerie au 30/09.")
ligne("cr_emis", "Créances ouvertes ≤ TTC facturé, par mois d'émission", "Créances", "Creances_Clients (TTC ouvert)", "Ventes_Detail (TTC facturé)",
      cr_mois, det_ttc, note="Écart = excédent du reste dû sur le facturé du mois (impossible s'il est > 0).", ecart_fn=lambda a, b: max(0.0, a - b))
ligne("cr_total", "Total créances recalculé vs affiché", "Créances", "Creances_Clients (Σ TTC)", "metrics.json",
      [None] * (NB - 1) + [sum(c["ttc"] for c in cr)], [None] * (NB - 1) + [M["creances"]["total"]], note="Balance âgée disponible au 30/09 uniquement.")

# ------------------------------------------------------------------ contrôles ponctuels
ponctuels = []


def pc(nom, ok, detail):
    ponctuels.append({"controle": nom, "ok": bool(ok), "detail": detail})


bad_px = [v for v in ventes if abs(v["qte"] * v["pu"] - v["ca"]) >= 0.01]
pc("CA HT = quantité × prix unitaire (chaque ligne)", not bad_px, f"{len(bad_px)} ligne(s) incohérente(s) sur {len(ventes)}")
neg = [v for v in ventes if v["ca"] <= 0 or v["qte"] <= 0]
pc("Aucun avoir ni ligne négative dans Ventes_Detail", not neg, f"{len(neg)} ligne(s) ≤ 0")
nums = sorted(int(v["num"].split("-")[1]) for v in ventes)
manquants = sorted(set(range(nums[0], nums[-1] + 1)) - set(nums))
pc("Numérotation des factures continue (pas de ligne manquante)", not manquants,
   f"FA-{nums[0]} à FA-{nums[-1]} : {len(set(nums))} numéros distincts, {len(manquants)} trou(s)" + (f" : {manquants[:10]}" if manquants else ""))
bornes = {}
for v in ventes:
    bornes.setdefault(v["m"], []).append(int(v["num"].split("-")[1]))
chevauch = [(m, m + 1) for m in range(1, NB) if m in bornes and m + 1 in bornes and max(bornes[m]) > min(bornes[m + 1])]
pc("Numéros de facture chronologiques (pas de facture glissée d'un mois à l'autre)", not chevauch,
   "plages mensuelles disjointes" if not chevauch else f"chevauchements : {chevauch}")
dates_hors = [v for v in ventes if v["date"].year != 2026 or v["m"] > NB]
pc("Toutes les factures datées de janv.–sept. 2026", not dates_hors, f"{len(dates_hors)} hors période")
cout_unit = defaultdict(set)
for v in ventes:
    cout_unit[v["produit"]].add(round(v["cout"] / v["qte"], 4))
var_cout = {p: s for p, s in cout_unit.items() if len(s) > 1}
pc("Coût de revient unitaire constant par produit", not var_cout, f"{len(var_cout)} produit(s) à coût variable")
fact_ttc = defaultdict(float)
vus = set()
for v in ventes:
    k = (v["num"], v["date"], v["produit"], v["qte"], v["ca"])
    if k in vus:
        continue  # une facture en double ne doit pas gonfler le plafond
    vus.add(k)
    fact_ttc[v["num"]] += v["ca"] * (1 + taux_tva(v["gamme"]))
inconnues = [c for c in cr if c["num"] not in fact_ttc]
pc("Chaque créance correspond à une facture du détail", not inconnues, f"{len(inconnues)} créance(s) sans facture")
sup = [c for c in cr if c["num"] in fact_ttc and c["ttc"] > fact_ttc[c["num"]] + SEUIL]
pc("Reste dû ≤ montant TTC de la facture d'origine", not sup, f"{len(sup)} créance(s) supérieure(s) à leur facture")
partiels = [c for c in cr if c["num"] in fact_ttc and c["ttc"] < fact_ttc[c["num"]] - SEUIL]
pc("Créances = factures entières (pas de règlement partiel)", True, f"{len(partiels)} règlement(s) partiel(s) — information, pas une erreur")
cli_fact = {v["num"]: v["client"] for v in ventes}
mauvais_cli = [c for c in cr if c["num"] in cli_fact and cli_fact[c["num"]] != c["client"]]
pc("Client de la créance = client de la facture", not mauvais_cli, f"{len(mauvais_cli)} incohérence(s)")
termes = defaultdict(set)
for c in cr:
    termes[c["client"]].add((c["echeance"] - c["date"]).days)
multi = {k: sorted(s) for k, s in termes.items() if len(s) > 1}
pc("Délai de paiement unique par client", not multi, "conditions : " + ", ".join(f"{k} {sorted(s)[0]} j" for k, s in sorted(termes.items())) if not multi else f"délais multiples : {multi}")
ws_c = wbc["Creances_Clients"]
diff_tr = 0
for row in range(2, ws_c.max_row + 1):
    ech = wb["Creances_Clients"].cell(row, 4).value
    if ech is None:
        continue
    j = max(0, (datetime(2026, 9, 30) - ech).days)
    tr = "Non échu" if j == 0 else "1-30 j" if j <= 30 else "31-60 j" if j <= 60 else "61-90 j" if j <= 90 else "> 90 j"
    if ws_c.cell(row, 6).value != j or ws_c.cell(row, 7).value != tr:
        diff_tr += 1
pc("Jours de retard et tranches (formules Excel) = recalcul", diff_tr == 0, f"{diff_tr} ligne(s) divergente(s)")
s60 = sum(c["ttc"] for c in cr if (datetime(2026, 9, 30) - c["echeance"]).days > 60)
pc("Créances > 60 j recalculées = affichées", abs(s60 - M["creances"]["plus_60"]["montant"]) < SEUIL, f"recalcul {e(s60)} vs affiché {e(M['creances']['plus_60']['montant'])}")
acomptes = sum(t_sort["Acomptes d'IS"])
ac_hyp = float(hyp["Acomptes d'IS versés 2026"])
pc("Acomptes d'IS (Tresorerie) = Hypotheses", abs(acomptes - ac_hyp) < SEUIL, f"{e(acomptes)} vs {e(ac_hyp)}")
loyer_ttc = reel["charges"]["Loyers & charges locatives"]
pc("Loyers décaissés = loyers P&L × 1,2 (TVA 20 %)", all(abs(t_sort["Loyers"][i] - loyer_ttc[i] * 1.2) < SEUIL for i in range(NB)), "contrôlé mois par mois")
creances_01_01 = sum(c["ttc"] for c in cr) - sum(det_ttc) + sum(t_enc)
pc("Créances implicites au 01/01 (balance 30/09 − TTC facturé + encaissements) positives", creances_01_01 > 0,
   f"{e(creances_01_01)} — cohérent si l'encours d'ouverture 2026 avait cet ordre de grandeur (non vérifiable : pas de balance d'ouverture)")

# ------------------------------------------------------------------ enquête sur les écarts
toutes_valeurs = []
for w in wb.worksheets:
    for row in w.iter_rows():
        for c in row:
            if isinstance(c.value, (int, float)):
                toutes_valeurs.append((w.title, c.coordinate, float(c.value)))


def chercher_valeur(x):
    return [f"{s}!{a}" for s, a, v in toutes_valeurs if abs(abs(v) - abs(x)) < 0.01]


causes = []


def enquete_ca(i, gamme=None):
    """Explique l'écart CA détail − P&L du mois i (0-based), éventuellement pour une gamme."""
    mois = i + 1
    lot = [v for v in ventes if v["m"] == mois and (gamme is None or v["gamme"] == gamme)]
    ecart = sum(v["ca"] for v in lot) - (reel["ca_g"][gamme][i] if gamme else reel["ca"][i])
    preuves, reste = [], ecart
    # 1. doublons
    groupes = defaultdict(list)
    for v in lot:
        groupes[v["num"]].append(v)
    for num, g in groupes.items():
        if len(g) > 1:
            ident = len({(x["date"], x["client"], x["produit"], x["qte"], x["ca"]) for x in g}) == 1
            exc = sum(x["ca"] for x in g[1:])
            reste -= exc
            preuves.append({"type": "doublon", "texte": f"{num} figure {len(g)} fois (lignes {', '.join(str(x['ligne']) for x in g)} de Ventes_Detail) : "
                            + ("lignes strictement identiques" if ident else "contenus différents") + f", {g[0]['client']}, {g[0]['produit']}, {e(g[0]['ca'])}.", "montant": round(exc, 2)})
    # 2. reste inexpliqué : signature
    sig = {}
    if abs(reste) >= SEUIL:
        cogs_gap = det_cout[i] - reel["cogs"][i]
        dup_cout = sum(x["cout"] for g in groupes.values() if len(g) > 1 for x in g[1:])
        cogs_res = cogs_gap - dup_cout
        prix_seul = abs(cogs_res) < SEUIL
        rond = abs(reste - round(reste, -2)) < 0.01
        une = [v for v in lot if abs(v["ca"] - abs(reste)) < SEUIL]
        deux = [(a, b) for a, b in itertools.combinations(lot, 2) if abs(a["ca"] + b["ca"] - abs(reste)) < SEUIL]
        ailleurs = [x for x in chercher_valeur(reste)]
        sig = {"prix_seul": prix_seul, "rond": rond}
        preuves.append({"type": "signature", "texte": (f"Reste {e(reste)} après doublons. Coût des ventes du mois : "
                        + ("identique entre détail et compta (écart de coût nul hors doublon) → l'écart porte uniquement sur le prix, pas sur des volumes."
                           if prix_seul else f"écart de coût {e(cogs_res)} hors doublon → des quantités sont en jeu.")), "montant": round(reste, 2)})
        preuves.append({"type": "recherche", "texte": f"Ligne unique de {e(abs(reste))} dans le détail : {len(une)} ; paire de lignes : {len(deux)} ; "
                        f"montant présent ailleurs dans le classeur : {', '.join(ailleurs) if ailleurs else 'nulle part'} ; montant rond : {'oui' if rond else 'non'}."})
        if reste > 0:
            preuves.append({"type": "recherche", "texte": "Aucun avoir (ligne négative) dans le détail, numérotation continue, aucune facture hors période : "
                            "ce n'est ni une ligne manquante ni un décalage de mois."})
    if not preuves:
        return None
    if abs(reste) < SEUIL:
        cause, cert = "Doublon de facture dans le détail des ventes : il explique 100 % de l'écart.", "établie"
    elif sig.get("prix_seul") and reste > 0:
        cause = (f"Avoir ou remise de {e(reste)} comptabilisé dans le CA mais absent du détail des factures. "
                 "La signature le désigne : effet prix pur (coût inchangé, donc pas de retour de marchandise), montant rond typique d'une écriture manuelle (RFA, coopération commerciale, ristourne).")
        cert = "probable"
        if gamme:
            cause += f" Écart concentré sur la gamme {gamme}."
    else:
        cause, cert = "Écart non expliqué par les données disponibles.", "inconnue"
    return {"mois": LIB[i], "perimetre": gamme or "Total", "ecart": round(ecart, 2), "cause": cause, "certitude": cert, "preuves": preuves}


for i in range(NB):
    if abs(det_ca[i] - reel["ca"][i]) >= SEUIL:
        gam_ko = [g for g in GAMMES if abs(det_g[g][i] - reel["ca_g"][g][i]) >= SEUIL]
        for g in gam_ko:
            causes.append(dict(enquete_ca(i, g), controle="ca"))
    if abs(det_cout[i] - reel["cogs"][i]) >= SEUIL:
        lot = [v for v in ventes if v["m"] == i + 1]
        cnt = defaultdict(list)
        for v in lot:
            cnt[v["num"]].append(v)
        dup = sum(x["cout"] for g in cnt.values() if len(g) > 1 for x in g[1:])
        gap = det_cout[i] - reel["cogs"][i]
        causes.append({"controle": "cogs", "mois": LIB[i], "perimetre": "Total", "ecart": round(gap, 2),
                       "certitude": "établie" if abs(gap - dup) < SEUIL else "inconnue",
                       "cause": ("Coût de revient de la facture en doublon compté deux fois dans le détail : explique 100 % de l'écart."
                                 if abs(gap - dup) < SEUIL else "Écart de coût non expliqué."),
                       "preuves": [{"type": "doublon", "texte": f"Coût des lignes en double : {e(dup)} ; écart constaté : {e(gap)}.", "montant": round(dup, 2)}]})
for L in lignes:
    if L["id"] in ("ca", "cogs") or L["id"].startswith("ca_") and L["id"] != "ca_aff":
        continue
    for i, c in enumerate(L["cellules"]):
        if c["ok"] is False:
            causes.append({"controle": L["id"], "mois": LIB[i], "perimetre": "Total", "ecart": c["ecart"], "certitude": "inconnue",
                           "cause": f"{L['controle']} : {L['a_label']} {e(c['a'])} vs {L['b_label']} {e(c['b'])}. À investiguer.", "preuves": []})

nb_rouge = sum(1 for L in lignes for c in L["cellules"] if c["ok"] is False)
nb_vert = sum(1 for L in lignes for c in L["cellules"] if c["ok"] is True)
synthese = [f"{nb_vert} cellules vertes, {nb_rouge} rouges sur {nb_vert + nb_rouge} contrôles mensuels ; "
            f"{sum(1 for p in ponctuels if p['ok'])}/{len(ponctuels)} contrôles ponctuels OK."]
aff_ko = [L for L in lignes if L["categorie"] == "Cockpit" and any(c["ok"] is False for c in L["cellules"])]
synthese.append("Les chiffres affichés par le cockpit (CA, EBITDA réel et budget) sont identiques au recalcul indépendant." if not aff_ko
                else "ATTENTION : le cockpit affiche des chiffres différents du recalcul indépendant : " + ", ".join(L["controle"] for L in aff_ko) + ".")
for c in causes:
    synthese.append(f"{c['mois']} — {c['perimetre']} ({e(c['ecart'])}) : {c['cause']} [{c['certitude']}]")

proof = {"genere_le": datetime.now().strftime("%d/%m/%Y %H:%M"), "source": XLSX.name, "metrics_genere_le": M["meta"]["genere_le"],
         "seuil": SEUIL, "mois": LIB, "lignes": lignes, "ponctuels": ponctuels, "causes": causes, "synthese": synthese,
         "methode": "verify.py relit l'Excel cellule par cellule (openpyxl), sans pandas et sans importer build_metrics.py. "
                    "Il ne prend que les cellules saisies, pas les formules, puis compare à l'Excel (valeurs de formules en cache) et à metrics.json (ce qu'affiche le cockpit)."}
OUT.write_text(json.dumps(proof, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
html = HTML.read_text(encoding="utf-8")
MARK = "/*__PROOF_JSON__*/null"
if MARK not in html:
    sys.exit("cockpit.html ne contient pas le marqueur Proof Center : relancez build_metrics.py avant verify.py.")
HTML.write_text(html.replace(MARK, json.dumps(proof, ensure_ascii=False, default=str).replace("</", "<\\/")), encoding="utf-8")

print("\n".join(synthese))
print("\nContrôles ponctuels :")
for p in ponctuels:
    print(f"  [{'OK' if p['ok'] else 'KO'}] {p['controle']} — {p['detail']}")
